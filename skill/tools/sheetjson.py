#!/usr/bin/env python3
"""Convert a CRM spreadsheet to JSON and back, so an agent can edit it as text.

    python3 skill/tools/sheetjson.py dump examples/career/crm.xlsx > draft.json
    python3 skill/tools/sheetjson.py load draft.json work/crm.xlsx

JSON shape: {"TabName": [["header1", "header2", ...], ["row1 value", ...], ...], ...}
Tabs keep their order. Every value is text; dates are written as YYYY-MM-DD text.
`load` styles the result for humans: bold, frozen header row and sensible column widths.
"""
from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def dump(xlsx: Path) -> dict:
    wb = load_workbook(xlsx, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = [[cell_text(v) for v in r] for r in ws.iter_rows(values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((len(r) - next((i for i, v in enumerate(reversed(r)) if v), len(r)) for r in rows), default=0)
        out[ws.title] = [r[:width] for r in rows]
    return out


def load(data: dict, xlsx: Path) -> None:
    if not isinstance(data, dict) or not data:
        sys.exit("the JSON has no tabs: expected {\"TabName\": [[header, ...], [row, ...]], ...}")
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E8F0EE")
    for name, rows in data.items():
        if len(name) > 31 or any(c in name for c in "[]:*?/\\"):
            sys.exit(f"tab name '{name}': at most 31 characters and none of []:*?/\\")
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(["" if v is None else str(v) for v in r])
        if not rows:
            continue
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head_fill
        ws.freeze_panes = "A2"
        for col in ws.columns:
            longest = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, longest + 2), 60)
            for c in col[1:]:
                c.alignment = Alignment(vertical="top", wrap_text=longest > 60)
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)


def main() -> None:
    if len(sys.argv) == 3 and sys.argv[1] == "dump":
        json.dump(dump(Path(sys.argv[2])), sys.stdout, ensure_ascii=False, indent=1)
        print()
    elif len(sys.argv) == 4 and sys.argv[1] == "load":
        load(json.loads(Path(sys.argv[2]).read_text()), Path(sys.argv[3]))
        print(f"wrote {sys.argv[3]}")
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
